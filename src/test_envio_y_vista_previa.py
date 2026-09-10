"""
Tests del flujo de envío: vista previa con marca de agua, «Descargar PDF y
marcar enviada», la bandeja de pendientes y la protección del Excel fiscal
frente al doble clic.

Uso (dentro de CT 104, con BD de prueba en /tmp):
    cd /mnt/empresa/proforma-admin/src && python3 -m pytest test_envio_y_vista_previa.py -v
"""
import base64
import os
import sys
import threading
import time

# Aislar de producción ANTES de importar db/app/excel (leen el env al importar).
TEST_DB = '/tmp/test_envio.db'
TEST_XLSX = '/tmp/test_envio.xlsx'
os.environ['DB_PATH'] = TEST_DB
os.environ['PDF_DIR'] = '/tmp/test_envio_pdf'
os.environ['EXCEL_PATH'] = TEST_XLSX
os.environ['EXCEL_BACKUP_DIR'] = '/tmp/test_envio_bak'
os.environ['EXCEL_PENDING_FILE'] = '/tmp/test_envio_pendientes.json'
os.environ['EXCEL_LOCK_FILE'] = '/tmp/test_envio.lock'

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pytest  # noqa: E402

import db  # noqa: E402
import excel  # noqa: E402

AUTH = {'Authorization': 'Basic ' + base64.b64encode(b'admin:admin').decode()}


@pytest.fixture(autouse=True)
def bd_limpia():
    db.DB_PATH = TEST_DB  # el módulo db se importa una vez por proceso
    for sufijo in ('', '-wal', '-shm'):
        try:
            os.remove(TEST_DB + sufijo)
        except FileNotFoundError:
            pass
    for f in (TEST_XLSX, os.environ['EXCEL_PENDING_FILE']):
        try:
            os.remove(f)
        except FileNotFoundError:
            pass
    for nombre in os.listdir(os.environ['PDF_DIR']) if os.path.isdir(os.environ['PDF_DIR']) else []:
        os.remove(os.path.join(os.environ['PDF_DIR'], nombre))
    db.init_db()
    db.set_serie_config('PRO', '{serie}-{anio}-{n}', 4)
    with db.get_db() as conn:
        conn.execute("INSERT INTO clientes (id, nombre_agencia, nif_cif) VALUES (1, 'Traditional Tours SLU', 'B11111111')")
        conn.execute("INSERT INTO cuentas (id, nombre, predeterminada) VALUES (1, 'Principal', 1)")
    yield


@pytest.fixture()
def client():
    from app import app
    app.config['TESTING'] = True
    with app.test_client() as c:
        yield c


def _crear_borrador(cliente_id=1):
    with db.get_db() as conn:
        conn.execute(
            """INSERT INTO proformas (numero_proforma, fecha, cliente_id, cuenta_id, estado,
               base, iva_total, total, total_suplidos, trimestre, numero_secuencial)
               VALUES ('PRO-2026-0001', '2026-09-10', ?, 1, 'borrador', 100, 21, 121, 121, 3, 1)""",
            (cliente_id,))
        pid = conn.execute("SELECT id FROM proformas").fetchone()['id']
        conn.execute(
            """INSERT INTO proforma_lineas (proforma_id, descripcion, cantidad, precio,
               porcentaje_iva, importe) VALUES (?, 'Visita guiada', 1, 100, 21, 100)""", (pid,))
    return pid


def _campo(pid, nombre):
    with db.get_db() as conn:
        return conn.execute(f"SELECT {nombre} FROM proformas WHERE id = ?", (pid,)).fetchone()[nombre]


def _filas_excel():
    if not os.path.exists(TEST_XLSX):
        return 0
    from openpyxl import load_workbook
    ws = load_workbook(TEST_XLSX).active
    return sum(1 for fila in ws.iter_rows(min_row=2) if fila[2].value)


# ── Vista previa ─────────────────────────────────────────────────────────────

def test_el_pdf_de_un_borrador_es_vista_previa_y_no_se_cachea(client):
    pid = _crear_borrador()
    r = client.get(f'/proformas/{pid}/pdf', headers=AUTH)
    assert r.status_code == 200
    assert r.data[:4] == b'%PDF'
    assert 'BORRADOR' in r.headers['Content-Disposition']
    # no debe dejar rastro: ni ruta_pdf ni fichero en el NAS
    assert _campo(pid, 'ruta_pdf') is None
    carpeta = os.environ['PDF_DIR']
    assert not os.path.isdir(carpeta) or not os.listdir(carpeta)


def test_descargar_la_vista_previa_sella_la_fecha_solo_la_primera_vez(client):
    pid = _crear_borrador()
    assert _campo(pid, 'pdf_previsualizado_en') is None
    client.get(f'/proformas/{pid}/pdf', headers=AUTH)
    primero = _campo(pid, 'pdf_previsualizado_en')
    assert primero is not None
    time.sleep(1.1)                      # el sello tiene resolución de segundos
    client.get(f'/proformas/{pid}/pdf', headers=AUTH)
    assert _campo(pid, 'pdf_previsualizado_en') == primero


def test_un_borrador_con_pdf_cacheado_sigue_dando_vista_previa(client):
    """Los siete PDF que quedaron en el NAS con el sello «Borrador» impreso no
    deben volver a servirse nunca como definitivos."""
    pid = _crear_borrador()
    os.makedirs(os.environ['PDF_DIR'], exist_ok=True)
    falso = os.path.join(os.environ['PDF_DIR'], 'PRO-2026-0001.pdf')
    with open(falso, 'wb') as f:
        f.write(b'%PDF-1.7 pdf viejo con sello Borrador')
    with db.get_db() as conn:
        conn.execute("UPDATE proformas SET ruta_pdf = ? WHERE id = ?", (falso, pid))

    r = client.get(f'/proformas/{pid}/pdf', headers=AUTH)
    assert 'BORRADOR' in r.headers['Content-Disposition']
    assert b'pdf viejo' not in r.data


# ── Descargar y marcar enviada ───────────────────────────────────────────────

def test_enviar_y_descargar_marca_enviada_y_redirige_a_la_descarga(client):
    pid = _crear_borrador()
    r = client.post(f'/proformas/{pid}/enviar-y-descargar', headers=AUTH)
    assert r.status_code == 302
    assert 'descargar=1' in r.headers['Location']
    assert _campo(pid, 'estado') == 'enviada'
    assert _campo(pid, 'ruta_pdf') is None      # se regenera ya como definitivo
    assert _filas_excel() == 1


def test_el_pdf_de_una_enviada_es_definitivo_y_se_cachea(client):
    pid = _crear_borrador()
    client.post(f'/proformas/{pid}/enviar-y-descargar', headers=AUTH)
    r = client.get(f'/proformas/{pid}/pdf', headers=AUTH)
    assert r.status_code == 200
    assert 'BORRADOR' not in r.headers['Content-Disposition']
    ruta = _campo(pid, 'ruta_pdf')
    assert ruta is not None and os.path.exists(ruta)


def test_enviar_y_descargar_sin_cliente_se_bloquea(client):
    """Sin cliente, la fila del Excel de Hacienda saldría sin agencia ni NIF."""
    pid = _crear_borrador(cliente_id=None)
    r = client.post(f'/proformas/{pid}/enviar-y-descargar', headers=AUTH)
    assert r.status_code == 302
    assert '/editar' in r.headers['Location']
    assert _campo(pid, 'estado') == 'borrador'
    assert _filas_excel() == 0


# ── El Excel de Hacienda no admite duplicados ────────────────────────────────

def test_doble_envio_secuencial_no_duplica_la_fila(client):
    pid = _crear_borrador()
    client.post(f'/proformas/{pid}/enviar', headers=AUTH)
    client.post(f'/proformas/{pid}/enviar', headers=AUTH)
    assert _filas_excel() == 1


def test_dos_envios_solapados_no_duplican_la_fila(client):
    """El servicio es multihilo: un doble clic mandaba dos peticiones a la vez y
    las dos escribían. Se fuerza el solape retrasando a la primera dentro del
    lock del Excel."""
    pid = _crear_borrador()
    original = excel._save_con_reintentos

    def lento(wb):
        time.sleep(0.6)
        return original(wb)

    excel._save_con_reintentos = lento
    try:
        hilos = [threading.Thread(target=lambda: excel.registrar_proforma(pid)),
                 threading.Thread(target=lambda: (time.sleep(0.2), excel.registrar_proforma(pid)))]
        for h in hilos:
            h.start()
        for h in hilos:
            h.join()
    finally:
        excel._save_con_reintentos = original

    assert _filas_excel() == 1


def test_enviar_una_cobrada_no_toca_el_excel(client):
    pid = _crear_borrador()
    client.post(f'/proformas/{pid}/enviar', headers=AUTH)
    with db.get_db() as conn:
        conn.execute("UPDATE proformas SET estado = 'cobrada' WHERE id = ?", (pid,))
    client.post(f'/proformas/{pid}/enviar', headers=AUTH)
    assert _filas_excel() == 1
    assert _campo(pid, 'estado') == 'cobrada'


# ── Deshacer envío ───────────────────────────────────────────────────────────

def test_desenviar_borra_el_pdf_definitivo(client):
    pid = _crear_borrador()
    client.post(f'/proformas/{pid}/enviar-y-descargar', headers=AUTH)
    client.get(f'/proformas/{pid}/pdf', headers=AUTH)
    ruta = _campo(pid, 'ruta_pdf')
    assert os.path.exists(ruta)

    client.post(f'/proformas/{pid}/desenviar', headers=AUTH)
    assert _campo(pid, 'estado') == 'borrador'
    assert _campo(pid, 'ruta_pdf') is None
    assert not os.path.exists(ruta)
    assert _filas_excel() == 0


# ── Bandeja de pendientes ────────────────────────────────────────────────────

def test_la_bandeja_lista_los_borradores_con_pdf_descargado(client):
    pid = _crear_borrador()
    assert b'PDF ya descargado' not in client.get('/proformas', headers=AUTH).data

    client.get(f'/proformas/{pid}/pdf', headers=AUTH)          # se descarga -> se sella
    html = client.get('/proformas', headers=AUTH).data.decode()
    assert 'PDF ya descargado' in html
    assert 'PRO-2026-0001' in html

    client.post(f'/proformas/{pid}/enviar', headers=AUTH)      # al enviarla desaparece
    assert 'PDF ya descargado' not in client.get('/proformas', headers=AUTH).data.decode()


def test_marcar_enviada_desde_la_bandeja_vuelve_al_listado(client):
    pid = _crear_borrador()
    client.get(f'/proformas/{pid}/pdf', headers=AUTH)
    r = client.post(f'/proformas/{pid}/enviar', headers=AUTH, data={'next': 'lista'})
    assert r.status_code == 302
    assert r.headers['Location'].endswith('/proformas')
    assert _campo(pid, 'estado') == 'enviada'


def test_la_migracion_sella_los_borradores_que_ya_tenian_pdf():
    """Los siete borradores legados: se les pone la fecha de su fichero PDF y no
    se les toca nada más."""
    pid = _crear_borrador()
    os.makedirs(os.environ['PDF_DIR'], exist_ok=True)
    ruta = os.path.join(os.environ['PDF_DIR'], 'legada.pdf')
    with open(ruta, 'wb') as f:
        f.write(b'%PDF')
    with db.get_db() as conn:
        conn.execute("UPDATE proformas SET ruta_pdf = ?, pdf_previsualizado_en = NULL WHERE id = ?",
                     (ruta, pid))
        conn.execute("ALTER TABLE proformas RENAME COLUMN pdf_previsualizado_en TO _viejo")

    with db.get_db() as conn:
        db._migrate_add_pdf_previsualizado_en(conn)

    assert _campo(pid, 'pdf_previsualizado_en') is not None
    assert _campo(pid, 'estado') == 'borrador'      # no la envía
    assert _campo(pid, 'exportada_excel') == 0      # no toca el Excel
